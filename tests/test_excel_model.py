import pandas as pd
from openpyxl import load_workbook

from supplypulse import excel_model


def _rec(n=5):
    return pd.DataFrame({
        "sku_id": [f"SKU-{i:03d}" for i in range(n)], "category": ["Snacks", "Pet"] * 3 if n > 5 else ["Snacks", "Pet", "Snacks", "Pet", "Snacks"],
        "abc_class": "A", "unit_cost": 5.0, "unit_price": 8.0, "lead_time_weeks": 2, "case_pack": 12, "moq": 24,
        "sigma_cover": 40.0, "z": 1.6, "safety_stock": 64.0, "expected_fill_rate": 0.99,
        "on_hand": 100.0, "on_order": 50.0, "weekly_forecast": 60.0, "cover_demand": 180.0,
        "order_up_to": 244.0, "inventory_position": 150.0, "order_qty": 96, "order_value": 480.0,
        "weeks_of_cover": 1.7, "status": "Order",
    })


def test_workbook_has_live_formulas_and_named_inputs(tmp_path, cfg):
    path = excel_model.build(_rec(), tmp_path / "m.xlsx", cfg)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Assumptions", "Planner", "How it works"]
    for name in ("Mode", "GlobalSL", "Uplift", "LTDelay", "HoldRate"):
        assert name in wb.defined_names
    ws = wb["Planner"]
    assert str(ws["U5"].value).startswith("=")          # order quantity is a formula
    assert "NORMSINV" in ws["O5"].value                  # global service level path
    assert wb["Summary"]["B5"].value.startswith("=SUM(")
    assert ws["A5"].value == "SKU-000" and ws.max_row == 4 + 5
