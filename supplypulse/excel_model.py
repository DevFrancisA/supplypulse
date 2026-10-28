"""Excel what-if replenishment model.

Everything a planner can change is an input cell on `Assumptions`; every number on
`Planner` and `Summary` is a live formula, so the workbook recalculates in Excel
(or LibreOffice) without Python.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, TEAL, INK, MUTED = "1F2A44", "0F766E", "1F2937", "6B7280"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
CALC_FILL = PatternFill("solid", fgColor="F3F6FA")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Planner layout ------------------------------------------------------------
HEADERS = [
    ("SKU", 11), ("Category", 14), ("ABC", 6), ("Unit cost", 10), ("Lead time (wks)", 11),
    ("Case pack", 9), ("MOQ", 8), ("Avg weekly forecast", 12), ("Sigma over cover (units)", 13),
    ("Optimised z", 10), ("On hand", 10), ("On order", 10),
    # --- calculated from here ---
    ("Weekly demand (scenario)", 13), ("Lead time (scenario)", 11), ("z used", 8),
    ("Demand over cover", 12), ("Sigma (scenario)", 11), ("Safety stock", 11),
    ("Order-up-to level", 12), ("Inventory position", 11), ("Order qty", 10),
    ("Order value", 12), ("Safety stock value", 13), ("Weeks of cover", 9),
    ("Exp. shortage / wk", 11), ("Status", 12), ("Annual SS holding cost", 13),
]
FIRST_CALC = 13  # column M
START_ROW = 5


def _style_header(ws, row: int, n: int):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEADER_FILL, Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def build(rec: pd.DataFrame, path: Path, cfg) -> Path:
    wb = Workbook()
    summary, assume, planner = wb.active, wb.create_sheet("Assumptions"), wb.create_sheet("Planner")
    summary.title = "Summary"
    notes = wb.create_sheet("How it works")

    # ---------------- Assumptions ----------------
    assume["A1"] = "What-if assumptions"
    assume["A1"].font = Font(bold=True, size=16, color=NAVY)
    assume["A2"] = "Yellow cells are inputs. Change them and Planner / Summary recalculate."
    assume["A2"].font = Font(italic=True, color=MUTED)
    rows = [
        ("Service level mode", "Optimised per SKU", "Mode", "Optimised per SKU uses the PuLP solution; Global applies one service level to every SKU."),
        ("Global cycle service level", 0.95, "GlobalSL", "Probability of not stocking out in a replenishment cycle (used when mode = Global)."),
        ("Demand uplift", 0.0, "Uplift", "Scenario change vs. forecast, e.g. 0.10 = +10% demand (scales mean and sigma)."),
        ("Lead-time delay (weeks)", 0, "LTDelay", "Extra weeks added to every supplier lead time."),
        ("Annual holding rate", cfg.holding_rate, "HoldRate", "Share of unit cost per year."),
    ]
    for i, (label, val, name, help_) in enumerate(rows, start=4):
        assume.cell(i, 1, label).font = Font(bold=True)
        c = assume.cell(i, 2, val)
        c.fill, c.border, c.font = INPUT_FILL, BOX, Font(color="0000FF", bold=True)
        c.alignment = Alignment(horizontal="center")
        assume.cell(i, 3, help_).font = Font(color=MUTED)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Assumptions!$B${i}")
    assume["B5"].number_format = assume["B6"].number_format = assume["B8"].number_format = "0.0%"
    dv = DataValidation(type="list", formula1='"Optimised per SKU,Global"', allow_blank=False)
    dv.add("B4"); assume.add_data_validation(dv)
    dv2 = DataValidation(type="decimal", operator="between", formula1="0.5", formula2="0.9999")
    dv2.add("B5"); assume.add_data_validation(dv2)
    dv3 = DataValidation(type="decimal", operator="between", formula1="-0.9", formula2="3")
    dv3.add("B6"); assume.add_data_validation(dv3)
    dv4 = DataValidation(type="whole", operator="between", formula1="0", formula2="12")
    dv4.add("B7"); assume.add_data_validation(dv4)
    assume.column_dimensions["A"].width = 30
    assume.column_dimensions["B"].width = 20
    assume.column_dimensions["C"].width = 95

    # ---------------- Planner ----------------
    planner["A1"] = "Replenishment planner"
    planner["A1"].font = Font(bold=True, size=16, color=NAVY)
    planner["A2"] = "Columns A-L: data from the Python pipeline. Columns M-AA: live formulas driven by the Assumptions sheet."
    planner["A2"].font = Font(italic=True, color=MUTED)
    for c, (h, w) in enumerate(HEADERS, start=1):
        planner.cell(4, c, h)
        planner.column_dimensions[get_column_letter(c)].width = w
    _style_header(planner, 4, len(HEADERS))
    planner.row_dimensions[4].height = 44

    n = len(rec)
    for i, r in enumerate(rec.itertuples(), start=START_ROW):
        avg_week = r.cover_demand / (r.lead_time_weeks + 1)
        vals = [r.sku_id, r.category, r.abc_class, r.unit_cost, r.lead_time_weeks, r.case_pack, r.moq,
                round(avg_week, 2), r.sigma_cover, r.z, r.on_hand, r.on_order]
        for c, v in enumerate(vals, start=1):
            cell = planner.cell(i, c, v)
            cell.border = BOX
        f = {
            13: f"=H{i}*(1+Uplift)",
            14: f"=E{i}+LTDelay",
            15: f'=IF(Mode="Global",NORMSINV(GlobalSL),J{i})',
            16: f"=M{i}*(N{i}+1)",
            17: f"=I{i}*(1+Uplift)*SQRT((N{i}+1)/(E{i}+1))",
            18: f"=O{i}*Q{i}",
            19: f"=P{i}+R{i}",
            20: f"=K{i}+L{i}",
            21: f"=IF(S{i}-T{i}<=0,0,MAX(G{i},CEILING(S{i}-T{i},F{i})))",
            22: f"=U{i}*D{i}",
            23: f"=R{i}*D{i}",
            24: f"=IF(M{i}>0,K{i}/M{i},0)",
            25: f"=Q{i}*(NORMDIST(O{i},0,1,FALSE)-O{i}*(1-NORMSDIST(O{i})))",
            26: f'=IF(U{i}>0,"Order",IF(T{i}>1.3*S{i},"Overstocked","Hold"))',
            27: f"=W{i}*HoldRate",
        }
        for c, formula in f.items():
            cell = planner.cell(i, c, formula)
            cell.fill, cell.border = CALC_FILL, BOX
    last = START_ROW + n - 1
    for col, fmt in {"D": "$#,##0.00", "H": "0.0", "I": "0.0", "J": "0.0", "K": "#,##0.0", "L": "#,##0.0",
                     "M": "0.0", "O": "0.00", "P": "#,##0.0", "Q": "0.0", "R": "#,##0.0", "S": "#,##0.0",
                     "T": "#,##0.0", "U": "#,##0", "V": "$#,##0", "W": "$#,##0", "X": "0.0", "Y": "0.00",
                     "AA": "$#,##0"}.items():
        for rr in range(START_ROW, last + 1):
            planner[f"{col}{rr}"].number_format = fmt
    planner.freeze_panes = planner.cell(START_ROW, 2)
    planner.auto_filter.ref = f"A4:{get_column_letter(len(HEADERS))}{last}"
    rng_status = f"Z{START_ROW}:Z{last}"
    planner.conditional_formatting.add(rng_status, CellIsRule(operator="equal", formula=['"Order"'],
                                       fill=PatternFill("solid", bgColor="FDE2E1"), font=Font(color="B42318", bold=True)))
    planner.conditional_formatting.add(rng_status, CellIsRule(operator="equal", formula=['"Overstocked"'],
                                       fill=PatternFill("solid", bgColor="FEF0C7"), font=Font(color="B54708", bold=True)))
    planner.conditional_formatting.add(rng_status, CellIsRule(operator="equal", formula=['"Hold"'],
                                       fill=PatternFill("solid", bgColor="D1FADF"), font=Font(color="067647", bold=True)))

    # ---------------- Summary ----------------
    P = lambda col: f"Planner!${col}${START_ROW}:${col}${last}"  # noqa: E731
    summary["A1"] = "SupplyPulse - replenishment summary"
    summary["A1"].font = Font(bold=True, size=18, color=NAVY)
    summary["A2"] = "All figures are live formulas over the Planner sheet. Change inputs on Assumptions."
    summary["A2"].font = Font(italic=True, color=MUTED)
    kpis = [
        ("SKUs to order this week", f'=COUNTIF({P("Z")},"Order")', "#,##0"),
        ("Total order value", f"=SUM({P('V')})", "$#,##0"),
        ("Units to order", f"=SUM({P('U')})", "#,##0"),
        ("Safety stock value", f"=SUM({P('W')})", "$#,##0"),
        ("Annual holding cost of safety stock", f"=SUM({P('AA')})", "$#,##0"),
        ("Expected fill rate (demand-weighted)", f"=1-SUM({P('Y')})/SUM({P('M')})", "0.0%"),
        ("SKUs flagged overstocked", f'=COUNTIF({P("Z")},"Overstocked")', "#,##0"),
    ]
    for i, (label, formula, fmt) in enumerate(kpis, start=4):
        summary.cell(i, 1, label).font = Font(bold=True, color=INK)
        c = summary.cell(i, 2, formula)
        c.number_format, c.font, c.border = fmt, Font(bold=True, size=13, color=TEAL), BOX
        c.alignment = Alignment(horizontal="right")
    summary["A13"] = "By category"
    summary["A13"].font = Font(bold=True, size=13, color=NAVY)
    for c, h in enumerate(["Category", "SKUs", "Order value", "Safety stock value", "SKUs to order"], start=1):
        summary.cell(14, c, h)
    _style_header(summary, 14, 5)
    cats = sorted(rec.category.unique())
    for i, cat in enumerate(cats, start=15):
        summary.cell(i, 1, cat)
        summary.cell(i, 2, f'=COUNTIF({P("B")},A{i})')
        summary.cell(i, 3, f'=SUMIF({P("B")},A{i},{P("V")})').number_format = "$#,##0"
        summary.cell(i, 4, f'=SUMIF({P("B")},A{i},{P("W")})').number_format = "$#,##0"
        summary.cell(i, 5, f'=COUNTIFS({P("B")},A{i},{P("Z")},"Order")')
        for c in range(1, 6):
            summary.cell(i, c).border = BOX
    end = 14 + len(cats)
    chart = BarChart()
    chart.type, chart.title, chart.style = "col", "Order value by category", 10
    chart.add_data(Reference(summary, min_col=3, min_row=14, max_row=end), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=15, max_row=end))
    chart.height, chart.width, chart.legend = 8, 16, None
    summary.add_chart(chart, "G3")
    summary.column_dimensions["A"].width = 38
    for col in "BCDE":
        summary.column_dimensions[col].width = 20

    # ---------------- Notes ----------------
    lines = [
        ("How the model works", True),
        ("Policy: weekly-review, order-up-to (base stock).  Order-up-to S = demand over the protection interval + safety stock.", False),
        ("Protection interval = lead time + 1 review week.  Safety stock = z x sigma, where sigma is the std of forecast error over that interval.", False),
        ("Order quantity = S - inventory position (on hand + on order), rounded up to the case pack and at least the MOQ.", False),
        ("Optimised z comes from a PuLP cost-minimisation (holding cost vs. lost-sales cost, subject to a fill-rate floor).", False),
        ("Switch Assumptions!B4 to Global to see a one-size-fits-all service level, then compare safety stock value and expected fill rate.", False),
        ("Lead-time delay scales sigma by sqrt((L'+1)/(L+1)); demand uplift scales both the mean and sigma.", False),
        ("Expected shortage per week = sigma x G(z), where G is the standard normal loss function.", False),
    ]
    for i, (t, bold) in enumerate(lines, start=1):
        notes.cell(i, 1, t).font = Font(bold=bold, size=14 if bold else 11, color=NAVY if bold else INK)
    notes.column_dimensions["A"].width = 140

    path = Path(path)
    wb.save(path)
    return path
